import html
import statistics
import sys
import tempfile
import uuid
from io import BytesIO
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.shared import Inches, Pt


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
V_NS = "urn:schemas-microsoft-com:vml"
O_NS = "urn:schemas-microsoft-com:office:office"


def pt(value: float) -> str:
    return f"{value:.2f}pt"


def twips(value: float) -> int:
    return max(1, int(round(value * 20)))


def escape_attr(value: str) -> str:
    return html.escape(value, quote=True)


def escape_text(value: str) -> str:
    return html.escape(value, quote=False)


def line_text(line: dict) -> str:
    text = "".join(span.get("text", "") for span in line.get("spans", []))
    return " ".join(text.split()) if text.strip().isascii() else text.strip()


def line_font_size(line: dict) -> float:
    sizes = [float(span.get("size", 10)) for span in line.get("spans", []) if span.get("text", "").strip()]
    if not sizes:
        return 10.5
    return max(5.0, min(24.0, statistics.median(sizes)))


def line_font_family(line: dict) -> str:
    fonts = [span.get("font", "") for span in line.get("spans", []) if span.get("text", "").strip()]
    joined = " ".join(fonts).lower()
    if "times" in joined:
        return "Times New Roman"
    if "hei" in joined or "bold" in joined:
        return "SimHei"
    if "fang" in joined or "kai" in joined:
        return "FangSong"
    return "SimSun"


def is_bold(line: dict) -> bool:
    return any("bold" in span.get("font", "").lower() for span in line.get("spans", []))


def textbox_text_xml(text: str) -> str:
    parts = text.split("\n")
    result = []
    for index, part in enumerate(parts):
        if index:
            result.append("<w:br/>")
        result.append(f'<w:t xml:space="preserve">{escape_text(part)}</w:t>')
    return "".join(result)


def build_textbox_xml(shape_id: str, x: float, y: float, width: float, height: float, text: str, font: str, size: float, bold: bool) -> str:
    size_half_points = int(round(size * 2))
    weight = '<w:b/>' if bold else ''
    # VML textboxes are widely accepted by Word and allow absolute positioning.
    return f'''
<w:pict xmlns:w="{W_NS}" xmlns:v="{V_NS}" xmlns:o="{O_NS}">
  <v:shape id="{escape_attr(shape_id)}" type="#_x0000_t202"
    style="position:absolute;margin-left:{pt(x)};margin-top:{pt(y)};width:{pt(width)};height:{pt(height)};z-index:2;mso-position-horizontal-relative:page;mso-position-vertical-relative:page"
    filled="f" stroked="f" o:allowincell="f">
    <v:textbox inset="0,0,0,0" style="mso-fit-shape-to-text:false">
      <w:txbxContent>
        <w:p>
          <w:pPr>
            <w:spacing w:before="0" w:after="0" w:line="{twips(size * 1.12)}" w:lineRule="exact"/>
          </w:pPr>
          <w:r>
            <w:rPr>
              <w:rFonts w:ascii="{escape_attr(font)}" w:hAnsi="{escape_attr(font)}" w:eastAsia="{escape_attr(font)}"/>
              <w:sz w:val="{size_half_points}"/>
              <w:szCs w:val="{size_half_points}"/>
              {weight}
            </w:rPr>
            {textbox_text_xml(text)}
          </w:r>
        </w:p>
      </w:txbxContent>
    </v:textbox>
  </v:shape>
</w:pict>
'''


def extract_text_blocks(page: fitz.Page) -> list[dict]:
    result = []
    text_dict = page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE)
    for block in text_dict.get("blocks", []):
        if block.get("type") != 0:
            continue
        block_lines = []
        line_bboxes = []
        sizes = []
        fonts = []
        bold = False
        for line in block.get("lines", []):
            text = line_text(line)
            if not text:
                continue
            x0, y0, x1, y1 = [float(v) for v in line.get("bbox", (0, 0, 0, 0))]
            if x1 <= x0 or y1 <= y0:
                continue
            block_lines.append(text)
            line_bboxes.append((x0, y0, x1, y1))
            sizes.append(line_font_size(line))
            fonts.append(line_font_family(line))
            bold = bold or is_bold(line)
        if not block_lines:
            continue
        x0 = min(b[0] for b in line_bboxes)
        y0 = min(b[1] for b in line_bboxes)
        x1 = max(b[2] for b in line_bboxes)
        y1 = max(b[3] for b in line_bboxes)
        font = "Times New Roman" if fonts.count("Times New Roman") > len(fonts) / 2 else "SimSun"
        result.append(
            {
                "text": "\n".join(block_lines),
                "bbox": (x0, y0, x1, y1),
                "redact_bboxes": line_bboxes,
                "font": font,
                "size": max(5.0, min(24.0, statistics.median(sizes))),
                "bold": bold,
            }
        )
    return result


def render_textless_page(source: Path, page_index: int, lines: list[dict], zoom: float) -> bytes:
    # Work on a one-page copy so redaction never mutates the source document.
    original = fitz.open(source)
    one_page = fitz.open()
    one_page.insert_pdf(original, from_page=page_index, to_page=page_index)
    original.close()
    page = one_page[0]
    for item in lines:
        for x0, y0, x1, y1 in item.get("redact_bboxes", [item["bbox"]]):
            rect = fitz.Rect(x0 - 0.8, y0 - 0.8, x1 + 0.8, y1 + 0.8)
            page.add_redact_annot(rect, fill=(1, 1, 1))
    page.apply_redactions(images=0, graphics=0, text=0)
    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
    data = pix.tobytes("png")
    one_page.close()
    return data


def convert(source: Path, target: Path) -> None:
    if not source.is_file():
        raise SystemExit(f"PDF not found: {source}")

    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        target.unlink()

    pdf = fitz.open(source)
    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(0)
    section.header_distance = section.footer_distance = Inches(0)

    zoom = 2.2
    for page_index, page in enumerate(pdf):
        rect = page.rect
        page_width_in = rect.width / 72
        page_height_in = rect.height / 72
        section.page_width = Inches(page_width_in)
        section.page_height = Inches(page_height_in)

        lines = extract_text_blocks(page)
        background = BytesIO(render_textless_page(source, page_index, lines, zoom))

        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = 0
        paragraph.paragraph_format.space_after = 0
        paragraph.paragraph_format.line_spacing = 1
        run = paragraph.add_run()
        run.add_picture(background, width=Inches(page_width_in), height=Inches(page_height_in))

        for idx, item in enumerate(lines):
            x0, y0, x1, y1 = item["bbox"]
            height = max(y1 - y0 + 4.0, item["size"] * 1.35)
            width = max(x1 - x0 + 4, 8)
            y = max(0, y0 - 1.5)
            xml = build_textbox_xml(
                f"zust_text_{page_index + 1}_{idx}_{uuid.uuid4().hex[:8]}",
                x0,
                y,
                width,
                height,
                item["text"],
                item["font"],
                item["size"],
                item["bold"],
            )
            run._r.append(parse_xml(xml))

    pdf.close()
    doc.save(target)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: pdf-object-docx.py SOURCE.pdf TARGET.docx")
    convert(Path(sys.argv[1]), Path(sys.argv[2]))
